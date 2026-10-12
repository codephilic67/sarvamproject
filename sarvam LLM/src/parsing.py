"""Document upload validation and text extraction (PDF, DOCX, TXT, Markdown)."""

from __future__ import annotations

import hashlib
import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO, Callable, Optional, Union

from src.models import DocumentRecord, PageText

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
MAX_FILE_SIZE_BYTES = 20 * 1024 * 1024  # 20 MB
MIME_BY_EXT = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
}


class ParseError(Exception):
    """Raised when document extraction fails."""


@dataclass
class ParsedDocument:
    document: DocumentRecord
    pages: list[PageText]


def validate_filename(filename: str) -> str:
    name = Path(filename).name
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ParseError(
            f"Unsupported file type '{ext}'. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )
    return name


def validate_size(size: int, max_bytes: int = MAX_FILE_SIZE_BYTES) -> None:
    if size <= 0:
        raise ParseError("File is empty (0 bytes).")
    if size > max_bytes:
        mb = max_bytes / (1024 * 1024)
        raise ParseError(f"File exceeds maximum size of {mb:.0f} MB.")


def content_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _normalize_whitespace(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(data: bytes) -> list[PageText]:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise ParseError("PyMuPDF (fitz) is required for PDF extraction.") from exc

    pages: list[PageText] = []
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ParseError(f"Failed to open PDF: {exc}") from exc

    try:
        if doc.page_count == 0:
            raise ParseError("PDF has no pages.")
        for i in range(doc.page_count):
            page = doc.load_page(i)
            text = _normalize_whitespace(page.get_text("text") or "")
            pages.append(PageText(page_number=i + 1, text=text))
    finally:
        doc.close()

    if not any(p.text for p in pages):
        raise ParseError(
            "PDF extraction produced no text. The file may be scanned/image-only "
            "or encrypted."
        )
    return pages


def extract_docx(data: bytes) -> list[PageText]:
    try:
        from docx import Document
        from io import BytesIO
    except ImportError as exc:
        raise ParseError("python-docx is required for DOCX extraction.") from exc

    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise ParseError(f"Failed to open DOCX: {exc}") from exc

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    # Tables
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    text = _normalize_whitespace("\n\n".join(paragraphs))
    if not text:
        raise ParseError("DOCX extraction produced no text.")
    # DOCX has no reliable page breaks; treat as single logical page.
    return [PageText(page_number=1, text=text)]


def extract_plain_text(data: bytes) -> list[PageText]:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ParseError("Could not decode text file with common encodings.")

    text = _normalize_whitespace(text)
    if not text:
        raise ParseError("Text document is empty after normalization.")
    return [PageText(page_number=1, text=text)]


def extract_pages(filename: str, data: bytes) -> list[PageText]:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_pdf(data)
    if ext == ".docx":
        return extract_docx(data)
    if ext in {".txt", ".md", ".markdown"}:
        return extract_plain_text(data)
    raise ParseError(f"Unsupported extension: {ext}")


def parse_document(
    filename: str,
    data: bytes,
    existing_hash_lookup: Optional[Callable] = None,
) -> ParsedDocument:
    """Validate, extract, and build a DocumentRecord with page provenance."""
    safe_name = validate_filename(filename)
    validate_size(len(data))
    ext = Path(safe_name).suffix.lower()
    file_hash = content_hash(data)

    is_duplicate = False
    duplicate_of = None
    if existing_hash_lookup:
        prior = existing_hash_lookup(file_hash)
        if prior is not None:
            is_duplicate = True
            duplicate_of = (
                prior.document_id if hasattr(prior, "document_id") else str(prior)
            )

    document_id = f"doc_{uuid.uuid4().hex[:12]}"
    try:
        pages = extract_pages(safe_name, data)
        char_count = sum(len(p.text) for p in pages)
        if char_count == 0:
            raise ParseError("Extracted document has zero characters.")
        status = "extracted"
        error_message = None
    except ParseError as exc:
        pages = []
        char_count = 0
        status = "failed"
        error_message = str(exc)

    doc = DocumentRecord(
        document_id=document_id,
        filename=safe_name,
        file_type=ext.lstrip("."),
        file_size=len(data),
        content_hash=file_hash,
        page_count=len(pages),
        char_count=char_count,
        status=status,
        error_message=error_message,
        is_duplicate=is_duplicate,
        duplicate_of=duplicate_of,
    )
    return ParsedDocument(document=doc, pages=pages)


def read_upload(file_obj: Union[BinaryIO, bytes], filename: str) -> ParsedDocument:
    if isinstance(file_obj, bytes):
        data = file_obj
    else:
        data = file_obj.read()
    return parse_document(filename, data)
